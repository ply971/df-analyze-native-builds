# Framework-agnostic logic shared by df-analyze's GUIs (web/streamlit_app.py
# and desktop/app.py). No `st.*` / `dpg.*` calls belong in this module --
# keeping option lists and CLI-arg-assembly here in one place means the two
# UIs can't silently drift apart.
from __future__ import annotations

import json
import math
import re
import shlex
from dataclasses import asdict, dataclass, field, fields, replace
from pathlib import Path
from typing import Any, Callable, Optional

import numpy as np
import pandas as pd

from df_analyze.gui_options import option_argv, validate_options

CLASSIFIERS = ["catboost", "dummy", "gandalf", "knn", "lgbm", "lr", "mlp", "rf", "sgd"]
DEFAULT_CLASSIFIERS = ["dummy", "knn", "lgbm", "sgd", "lr"]
REGRESSORS = [
    "catboost",
    "dummy",
    "elastic",
    "gandalf",
    "knn",
    "lgbm",
    "mlp",
    "rf",
    "sgd",
]
DEFAULT_REGRESSORS = ["dummy", "knn", "lgbm", "sgd", "elastic"]
FEAT_SELECT_METHODS = ["filter", "embed", "wrap", "none"]
EMBED_MODELS = ["linear", "lgbm", "none"]
WRAPPER_METHODS = ["step-up", "step-down", "none"]
WRAPPER_MODELS = ["linear", "lgbm", "knn", "none"]
NORM_METHODS = ["robust", "minmax"]
NAN_METHODS = ["mean", "median", "impute"]
CLS_METRICS = ["acc", "sens", "spec", "ppv", "npv", "f1", "bal-acc"]
REG_METRICS = ["mae", "msqe", "mdae", "r2", "var-exp"]
ERROR_METRICS = {"mae", "msqe", "mdae"}  # lower is better, for sort direction

# Filter-based feature selection: which statistic measures a feature's
# association with the target (or which score a quick predictive fit is
# judged on), separately for continuous vs. categorical features and for
# classification vs. regression -- the CLI has always supported choosing
# these, the GUI just didn't expose them yet. ("relief" is a real
# --filter-method choice but is documented as CURRENTLY UNIMPLEMENTED in the
# CLI's own --help, so it's deliberately left out here.)
FILTER_METHODS = ["assoc", "pred"]
CONT_CLS_STATS = ["t", "U", "W", "corr", "cohen_d", "AUROC", "mut_info"]
CAT_CLS_STATS = ["mut_info", "H", "cramer_v"]
CONT_REG_STATS = ["pearson_r", "spearman_r", "mut_info", "F"]
CAT_REG_STATS = ["mut_info", "H"]
CLS_PRED_SCORES = ["acc", "auroc", "sens", "spec"]
REG_PRED_SCORES = ["mae", "msqe", "mdae", "r2", "var-exp"]

# Multi-target analysis (--targets/--mt-agg-strategy/--mt-top-k): the CLI has
# supported this since PR #52/#53 -- these are just its two aggregation
# choices, unrelated to feature selection above.
MT_AGG_STRATEGIES = ["borda", "freq"]

SEQUENTIAL_BLUE = "#2a78d6"
SUPPORTED_SUFFIXES = {".csv", ".parquet", ".xlsx", ".json"}
INPUT_MODES = ["Data table", "Annotated spreadsheet", "Training + test files"]
DELIMITERS = {"Comma": ",", "Tab": "\t", "Semicolon": ";", "Pipe": "|", "Space": " "}
PRESET_NAMES = ["Quick baseline", "Balanced", "Thorough"]

# Matches the CLI's own fallback when --outdir is omitted (see
# src/df_analyze/saving.py: `outdir = root or Path.home().resolve() / "df-analyze-outputs"`),
# rather than a repo-relative path that only makes sense for a dev-run web app.
DEFAULT_OUTDIR = Path.home() / "df-analyze-outputs"


@dataclass
class RunConfig:
    """Everything needed to build a df-analyze CLI invocation. Both GUIs build
    one of these from their widget state, then pass it to build_argv()."""

    data_path: Path
    target: str
    mode: str  # "classify" | "regress"
    models: list[str]
    feat_select: list[str]
    embed_model: str = "linear"
    wrapper_method: str = "step-up"
    wrapper_model: str = "linear"
    norm: str = "robust"
    nan: str = "median"
    htune_trials: int = 100
    htune_metric: str = "acc"
    test_val_size: float = 0.4
    seed: int = 0  # 0 => omit --seed (let the CLI use its own default)
    outdir: Path = field(default_factory=lambda: DEFAULT_OUTDIR)
    extra_args: str = ""
    categoricals: list[str] = field(default_factory=list)
    ordinals: list[str] = field(default_factory=list)
    drops: list[str] = field(default_factory=list)
    grouper: Optional[str] = None
    no_preds: bool = False
    adaptive_error: bool = False
    # Applied by the GUI itself before the CLI ever sees the data: rows are
    # deduplicated and the result written to a real, kept file in outdir
    # (not a throwaway temp file), so the analysis and its reproducible
    # command both refer to the exact data that was actually used.
    drop_duplicates: bool = False
    # Filter-based feature selection tuning (see FILTER_METHODS etc. above).
    filter_method: str = "assoc"
    filter_assoc_cont_classify: str = "mut_info"
    filter_assoc_cat_classify: str = "mut_info"
    filter_assoc_cont_regress: str = "F"
    filter_assoc_cat_regress: str = "H"
    filter_pred_classify: str = "acc"
    filter_pred_regress: str = "mae"
    # Kept as strings (not float) because the CLI itself accepts either an
    # integer feature count ("20") or a fraction of all features ("0.5") for
    # both of these -- forcing one numeric type here would silently drop the
    # other, valid form.
    n_feat_filter: str = "0.5"
    n_feat_wrapper: str = "20"
    # "Extra greedy" wrapper selection: only meaningful together with
    # feat_select including "wrap".
    redundant_wrapper_selection: bool = False
    redundant_threshold: float = 0.005
    redundant_corr_threshold: float = 0.8
    # A chainable, no-code data-manipulation pipeline (see TransformStep
    # below). Applied to the FULL dataset, in order, after drop_duplicates'
    # legacy behavior and before the CLI ever sees the data -- same
    # "GUI-side transform, real kept file, spliced into the run" approach.
    pipeline_steps: list[TransformStep] = field(default_factory=list)
    # Multi-target analysis: `target` above stays the primary target;
    # non-empty extra_targets switches build_argv() to the CLI's --targets
    # flag (target + extra_targets, comma-separated) instead of --target.
    # Empty extra_targets keeps argv byte-identical to single-target runs.
    extra_targets: list[str] = field(default_factory=list)
    mt_agg_strategy: str = "borda"
    # Kept as a string (not int) so "" cleanly means "omit --mt-top-k",
    # matching the existing n_feat_filter/n_feat_wrapper convention.
    mt_top_k: str = ""
    input_mode: str = "Data table"
    separator: str = ","
    test_files: list[str] = field(default_factory=list)
    tests_method: str = "list"
    advanced_options: dict[str, Any] = field(default_factory=dict)


def parse_extra_args(extra_args: str) -> list[str]:
    return shlex.split(extra_args) if extra_args.strip() else []


def build_argv(cfg: RunConfig) -> list[str]:
    """Build the df-analyze CLI argv -- WITHOUT a leading program name or
    interpreter. Callers prepend their own: the Streamlit app prepends
    [sys.executable, str(DF_ANALYZE)] and shells out; the desktop app
    prepends ["df-analyze.py"] onto sys.argv and calls the pipeline
    in-process (subprocess doesn't survive being frozen by PyInstaller)."""
    argv: list[str] = []
    input_flag = {"Data table": "--df", "Annotated spreadsheet": "--spreadsheet",
                  "Training + test files": "--df-train"}[cfg.input_mode]
    argv += [input_flag, str(cfg.data_path)]
    if cfg.separator != ",":
        argv += ["--separator", "tab" if cfg.separator == "\t" else cfg.separator]
    if cfg.test_files:
        argv += ["--df-tests", ",".join(cfg.test_files), "--df-tests-method", cfg.tests_method]
    if cfg.extra_targets:
        # Multi-target: --targets takes priority over --target in the CLI's
        # own reconciliation, so only one of the two flags is ever emitted.
        # De-duplicated defensively -- validate_config() is the real guard
        # against the target also appearing in extra_targets.
        targets = list(dict.fromkeys([cfg.target, *cfg.extra_targets]))
        argv += ["--targets", ",".join(targets)]
        if cfg.mt_agg_strategy != "borda":
            argv += ["--mt-agg-strategy", cfg.mt_agg_strategy]
        if cfg.mt_top_k.strip():
            argv += ["--mt-top-k", cfg.mt_top_k.strip()]
    else:
        argv += ["--target", cfg.target]
    argv += ["--mode", cfg.mode]
    argv += ["--outdir", str(cfg.outdir)]
    if cfg.mode == "classify":
        argv += ["--classifiers", *cfg.models]
        argv += ["--htune-cls-metric", cfg.htune_metric]
    else:
        argv += ["--regressors", *cfg.models]
        argv += ["--htune-reg-metric", cfg.htune_metric]
    argv += ["--feat-select", *cfg.feat_select]
    if "embed" in cfg.feat_select:
        argv += ["--embed-select", cfg.embed_model]
    else:
        argv += ["--embed-select", "none"]
    if "wrap" in cfg.feat_select:
        argv += ["--wrapper-select", cfg.wrapper_method]
        argv += ["--wrapper-model", cfg.wrapper_model]
        argv += ["--n-feat-wrapper", str(cfg.n_feat_wrapper)]
        if cfg.redundant_wrapper_selection:
            argv += ["--redundant-wrapper-selection"]
            argv += ["--redundant-threshold", str(cfg.redundant_threshold)]
            argv += ["--redundant-corr-threshold", str(cfg.redundant_corr_threshold)]
    else:
        argv += ["--wrapper-select", "none"]
    if "filter" in cfg.feat_select:
        argv += ["--filter-method", cfg.filter_method]
        if cfg.mode == "classify":
            argv += ["--filter-assoc-cont-classify", cfg.filter_assoc_cont_classify]
            argv += ["--filter-assoc-cat-classify", cfg.filter_assoc_cat_classify]
            argv += ["--filter-pred-classify", cfg.filter_pred_classify]
        else:
            argv += ["--filter-assoc-cont-regress", cfg.filter_assoc_cont_regress]
            argv += ["--filter-assoc-cat-regress", cfg.filter_assoc_cat_regress]
            argv += ["--filter-pred-regress", cfg.filter_pred_regress]
        argv += ["--n-feat-filter", str(cfg.n_feat_filter)]
    argv += ["--norm", cfg.norm]
    argv += ["--nan", cfg.nan]
    argv += ["--htune-trials", str(cfg.htune_trials)]
    argv += ["--test-val-size", str(cfg.test_val_size)]
    if cfg.seed:
        argv += ["--seed", str(cfg.seed)]
    for flag, columns in (
        ("--categoricals", cfg.categoricals),
        ("--ordinals", cfg.ordinals),
        ("--drops", cfg.drops),
    ):
        if columns:
            argv += [flag, ",".join(columns)]
    if cfg.grouper:
        argv += ["--grouper", cfg.grouper]
    if cfg.no_preds:
        argv += ["--no-preds"]
    if cfg.adaptive_error:
        argv += ["--adaptive-error"]
    # Typed advanced controls can replace the corresponding basic control.
    # Only explicitly enabled settings are emitted; old commands stay identical.
    advanced = option_argv(cfg.advanced_options)
    for flag in cfg.advanced_options:
        if flag in argv:
            start = argv.index(flag)
            end = start + 1
            while end < len(argv) and not argv[end].startswith("--"):
                end += 1
            del argv[start:end]
    argv += advanced
    argv += parse_extra_args(cfg.extra_args)
    return argv


def load_columns(path: Path) -> list[str]:
    """Peek at a data file's column names without loading the whole thing."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return list(pd.read_csv(path, nrows=5).columns)
    if suffix == ".parquet":
        import pyarrow.parquet as pq

        return list(pq.ParquetFile(path).schema.names)
    if suffix == ".xlsx":
        return list(pd.read_excel(path, nrows=5).columns)
    if suffix == ".json":
        return list(pd.read_json(path).columns)
    raise ValueError(f"Unsupported file type: {suffix}")


def load_preview(path: Path, max_rows: int = 2000, separator: str = ",", input_mode: str = "Data table") -> pd.DataFrame:
    """Load a bounded preview; JSON requires parsing the document first.

    CSV/Excel read only the requested rows. Parquet reads one bounded batch.
    Profile statistics always describe this preview, not the entire dataset.
    """
    if max_rows < 1:
        raise ValueError("Preview row limit must be positive.")
    if input_mode == "Annotated spreadsheet":
        from df_analyze.loading import load_spreadsheet

        return load_spreadsheet(path, separator)[0].head(max_rows)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return pd.read_csv(path, nrows=max_rows, sep=separator)
    if suffix == ".xlsx":
        return pd.read_excel(path, nrows=max_rows)
    if suffix == ".json":
        return pd.read_json(path).head(max_rows)
    if suffix == ".parquet":
        import pyarrow.parquet as pq

        parquet = pq.ParquetFile(path)
        batch = next(parquet.iter_batches(batch_size=max_rows), None)
        return (
            batch.to_pandas()
            if batch is not None
            else pd.DataFrame(columns=parquet.schema_arrow.names)
        )
    raise ValueError(f"Unsupported file type: {suffix}. Use CSV, Parquet, XLSX or JSON.")


def load_full_dataset(path: Path, separator: str = ",", input_mode: str = "Data table") -> pd.DataFrame:
    """Load the complete dataset, no row cap.

    Unlike load_preview (bounded, for fast exploration), a cleaning
    transformation like dropping duplicate rows has to see every row, not
    just the preview sample -- otherwise "cleaning" would silently only
    apply to whatever happened to be in the first couple thousand rows.
    """
    if input_mode == "Annotated spreadsheet":
        from df_analyze.loading import load_spreadsheet

        return load_spreadsheet(path, separator)[0]
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return pd.read_csv(path, sep=separator)
    if suffix == ".xlsx":
        return pd.read_excel(path)
    if suffix == ".json":
        return pd.read_json(path)
    if suffix == ".parquet":
        return pd.read_parquet(path)
    raise ValueError(f"Unsupported file type: {suffix}. Use CSV, Parquet, XLSX or JSON.")


def save_dataset(df: pd.DataFrame, path: Path) -> None:
    """Write a dataframe back out in whatever format its own suffix implies."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        df.to_csv(path, index=False)
    elif suffix == ".xlsx":
        df.to_excel(path, index=False)
    elif suffix == ".json":
        df.to_json(path, orient="records")
    elif suffix == ".parquet":
        df.to_parquet(path, index=False)
    else:
        raise ValueError(f"Unsupported file type: {suffix}. Use CSV, Parquet, XLSX or JSON.")


def drop_duplicate_rows(df: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    """Return (deduplicated frame, number of rows removed)."""
    cleaned = df.drop_duplicates()
    return cleaned, len(df) - len(cleaned)


# ---------------------------------------------------------------------------
# No-code data manipulation pipeline.
#
# A fixed, allow-listed catalog of real pandas/numpy operations, each exposed
# in the GUI as a dropdown + simple inputs (never a code/expression box). A
# TransformStep is pure, JSON-safe data -- never a callable or a code string
# -- so a saved/imported experiment JSON can't smuggle in arbitrary logic,
# the same stance config_from_json already takes with its allow-lists below.
#
# Every op function has the shape (df, **params) -> (new_df, summary): it
# never mutates its input, and it returns a short, human-readable sentence
# describing what changed, for the pipeline history / run log.
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class TransformOpSpec:
    """Registry metadata for one pipeline operation: its GUI display label
    and which category tab it belongs under. Not a parameter schema -- each
    op function below validates its own arguments."""

    label: str
    category: str  # "Column" | "Row" | "Missing" | "Feature"


@dataclass(frozen=True)
class TransformStep:
    """One pipeline step. `op` must be a key of TRANSFORM_OPS; `params` may
    only ever hold JSON-safe values (str/int/float/bool/list[str])."""

    op: str
    params: dict[str, Any] = field(default_factory=dict)


CAST_DTYPES = ["int64", "float64", "string", "category", "bool"]
TEXT_CLEAN_MODES = [
    "Trim",
    "Lowercase",
    "Uppercase",
    "Title case",
    "Remove punctuation",
    "Collapse whitespace",
]
FILL_STRATEGIES = [
    "Mean",
    "Median",
    "Mode",
    "Zero",
    "Custom value",
    "Forward fill",
    "Back fill",
]
FILTER_OPERATORS = ["==", "!=", ">", ">=", "<", "<=", "contains", "is missing", "is not missing"]
THRESHOLD_OPERATORS = ["==", "!=", ">", ">=", "<", "<="]
DATETIME_PARTS = ["year", "month", "day", "dayofweek", "quarter", "hour"]
COMBINE_OPERATORS = ["+", "-", "*", "/"]
CLIP_METHODS = ["Percentile", "Std devs from mean", "Manual bounds"]
DEDUPE_KEEP = ["first", "last", "none"]

TRANSFORM_OPS: dict[str, TransformOpSpec] = {
    "rename_column": TransformOpSpec("Rename column", "Column"),
    "cast_column_dtype": TransformOpSpec("Change data type", "Column"),
    "drop_columns": TransformOpSpec("Drop column(s)", "Column"),
    "clean_text_column": TransformOpSpec("Clean up text", "Column"),
    "extract_datetime_parts": TransformOpSpec("Extract date/time parts", "Column"),
    "filter_rows": TransformOpSpec("Filter rows", "Row"),
    "dedupe_rows": TransformOpSpec("Remove duplicate rows", "Row"),
    "drop_missing_rows": TransformOpSpec("Drop rows with missing values", "Row"),
    "fill_missing_column": TransformOpSpec("Fill missing values", "Missing"),
    "flag_missing_column": TransformOpSpec("Flag missing values", "Missing"),
    "clip_column_outliers": TransformOpSpec("Clip outliers", "Feature"),
    "threshold_flag_column": TransformOpSpec("Create 0/1 flag column", "Feature"),
    "combine_columns": TransformOpSpec("Combine two columns", "Feature"),
}


def _require_column(df: pd.DataFrame, column: str) -> None:
    if column not in df.columns:
        raise ValueError(f"Column '{column}' does not exist.")


def _require_new_column_name(df: pd.DataFrame, new_column: str) -> str:
    new_column = new_column.strip()
    if not new_column:
        raise ValueError("New column name cannot be blank.")
    if new_column in df.columns:
        raise ValueError(f"A column named '{new_column}' already exists.")
    return new_column


def rename_column(df: pd.DataFrame, column: str, new_name: str) -> tuple[pd.DataFrame, str]:
    _require_column(df, column)
    new_name = new_name.strip()
    if not new_name:
        raise ValueError("New column name cannot be blank.")
    if new_name != column and new_name in df.columns:
        raise ValueError(f"A column named '{new_name}' already exists.")
    result = df.rename(columns={column: new_name})
    return result, f"Renamed column '{column}' to '{new_name}'."


def cast_column_dtype(df: pd.DataFrame, column: str, dtype: str) -> tuple[pd.DataFrame, str]:
    _require_column(df, column)
    if dtype not in CAST_DTYPES:
        raise ValueError(f"Unsupported data type: {dtype}")
    result = df.copy()
    before = str(result[column].dtype)
    try:
        result[column] = result[column].astype(str) if dtype == "string" else result[column].astype(dtype)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"Could not cast '{column}' to {dtype}: {exc}") from exc
    return result, f"Cast '{column}' from {before} to {dtype}."


def drop_columns(df: pd.DataFrame, columns: list[str]) -> tuple[pd.DataFrame, str]:
    if not columns:
        raise ValueError("Choose at least one column to drop.")
    missing = [c for c in columns if c not in df.columns]
    if missing:
        raise ValueError("Column(s) not found: " + ", ".join(missing))
    result = df.drop(columns=columns)
    return result, "Dropped column(s): " + ", ".join(f"'{c}'" for c in columns) + "."


def clean_text_column(df: pd.DataFrame, column: str, mode: str) -> tuple[pd.DataFrame, str]:
    _require_column(df, column)
    if mode not in TEXT_CLEAN_MODES:
        raise ValueError(f"Unsupported text-cleaning mode: {mode}")
    result = df.copy()
    series = result[column].astype(str)
    if mode == "Trim":
        result[column], verb = series.str.strip(), "Trimmed whitespace from"
    elif mode == "Lowercase":
        result[column], verb = series.str.lower(), "Lowercased"
    elif mode == "Uppercase":
        result[column], verb = series.str.upper(), "Uppercased"
    elif mode == "Title case":
        result[column], verb = series.str.title(), "Title-cased"
    elif mode == "Remove punctuation":
        result[column] = series.str.replace(r"[^\w\s]", "", regex=True)
        verb = "Removed punctuation from"
    else:  # Collapse whitespace
        result[column] = series.str.replace(r"\s+", " ", regex=True).str.strip()
        verb = "Collapsed repeated whitespace in"
    return result, f"{verb} all values in '{column}'."


def extract_datetime_parts(df: pd.DataFrame, column: str, parts: list[str]) -> tuple[pd.DataFrame, str]:
    _require_column(df, column)
    if not parts:
        raise ValueError("Choose at least one date/time part to extract.")
    unknown = [p for p in parts if p not in DATETIME_PARTS]
    if unknown:
        raise ValueError("Unsupported date/time part(s): " + ", ".join(unknown))
    result = df.copy()
    parsed = pd.to_datetime(result[column], errors="coerce")
    if parsed.notna().sum() == 0:
        raise ValueError(f"Could not parse any values in '{column}' as dates/times.")
    new_columns = []
    for part in parts:
        new_col = _require_new_column_name(result, f"{column}_{part}")
        result[new_col] = getattr(parsed.dt, part)
        new_columns.append(new_col)
    return (
        result,
        f"Extracted {', '.join(parts)} from '{column}' into {len(new_columns)} new column(s).",
    )


def filter_rows(
    df: pd.DataFrame, column: str, operator: str, value: str, keep: bool = True
) -> tuple[pd.DataFrame, str]:
    _require_column(df, column)
    if operator not in FILTER_OPERATORS:
        raise ValueError(f"Unsupported comparison: {operator}")
    series = df[column]
    if operator == "is missing":
        mask = series.isna()
    elif operator == "is not missing":
        mask = series.notna()
    elif operator == "contains":
        mask = series.astype(str).str.contains(value, case=False, regex=False, na=False)
    else:
        numeric_series = pd.to_numeric(series, errors="coerce")
        try:
            numeric_value: Optional[float] = float(value)
        except (TypeError, ValueError):
            numeric_value = None
        if numeric_value is not None and numeric_series.notna().sum() > 0:
            compare_series, compare_value = numeric_series, numeric_value
        else:
            compare_series, compare_value = series.astype(str), value
        if operator == "==":
            mask = compare_series == compare_value
        elif operator == "!=":
            mask = compare_series != compare_value
        elif operator == ">":
            mask = compare_series > compare_value
        elif operator == ">=":
            mask = compare_series >= compare_value
        elif operator == "<":
            mask = compare_series < compare_value
        else:  # "<="
            mask = compare_series <= compare_value
    mask = mask.fillna(False)
    final_mask = mask if keep else ~mask
    result = df.loc[final_mask].reset_index(drop=True)
    removed = len(df) - len(result)
    verb = "Kept" if keep else "Removed"
    condition = (
        f"'{column}' {operator}"
        if operator in ("is missing", "is not missing")
        else f"'{column}' {operator} {value}"
    )
    return result, f"{verb} rows where {condition} ({removed:,} row(s) removed)."


def dedupe_rows(
    df: pd.DataFrame, subset: Optional[list[str]] = None, keep: str = "first"
) -> tuple[pd.DataFrame, str]:
    if keep not in DEDUPE_KEEP:
        raise ValueError(f"Unsupported keep option: {keep}")
    if subset:
        missing = [c for c in subset if c not in df.columns]
        if missing:
            raise ValueError("Column(s) not found: " + ", ".join(missing))
    keep_arg: Any = False if keep == "none" else keep
    result = df.drop_duplicates(subset=subset or None, keep=keep_arg).reset_index(drop=True)
    removed = len(df) - len(result)
    subset_desc = f" (subset: {', '.join(subset)})" if subset else ""
    keep_desc = "keeping none" if keep == "none" else f"keeping {keep}"
    return result, f"Removed {removed:,} duplicate row(s){subset_desc}, {keep_desc}."


def drop_missing_rows(df: pd.DataFrame, column: Optional[str] = None) -> tuple[pd.DataFrame, str]:
    if column and column not in df.columns:
        raise ValueError(f"Column '{column}' does not exist.")
    subset = [column] if column else None
    result = df.dropna(subset=subset).reset_index(drop=True)
    removed = len(df) - len(result)
    where = f"missing '{column}'" if column else "missing any value"
    return result, f"Dropped {removed:,} row(s) {where}."


def fill_missing_column(
    df: pd.DataFrame, column: str, strategy: str, custom_value: str = ""
) -> tuple[pd.DataFrame, str]:
    _require_column(df, column)
    if strategy not in FILL_STRATEGIES:
        raise ValueError(f"Unsupported fill strategy: {strategy}")
    result = df.copy()
    series = result[column]
    missing_before = int(series.isna().sum())
    if missing_before == 0:
        return result, f"No missing values in '{column}'; nothing to fill."
    if strategy == "Mean":
        fill_value = pd.to_numeric(series, errors="coerce").mean()
        result[column], desc = series.fillna(fill_value), f"the column mean ({fill_value:.4g})"
    elif strategy == "Median":
        fill_value = pd.to_numeric(series, errors="coerce").median()
        result[column], desc = series.fillna(fill_value), f"the column median ({fill_value:.4g})"
    elif strategy == "Mode":
        modes = series.mode(dropna=True)
        if modes.empty:
            raise ValueError(f"Column '{column}' has no non-missing values to compute a mode from.")
        fill_value = modes.iloc[0]
        result[column], desc = series.fillna(fill_value), f"the most common value ({fill_value!r})"
    elif strategy == "Zero":
        result[column], desc = series.fillna(0), "zero"
    elif strategy == "Custom value":
        if custom_value == "":
            raise ValueError("Enter a custom value to fill with.")
        try:
            parsed_value: Any = float(custom_value)
        except ValueError:
            parsed_value = custom_value
        result[column], desc = series.fillna(parsed_value), f"'{custom_value}'"
    elif strategy == "Forward fill":
        result[column], desc = series.ffill(), "the previous row's value (forward fill)"
    else:  # Back fill
        result[column], desc = series.bfill(), "the next row's value (back fill)"
    return result, f"Filled {missing_before:,} missing value(s) in '{column}' with {desc}."


def flag_missing_column(df: pd.DataFrame, column: str) -> tuple[pd.DataFrame, str]:
    _require_column(df, column)
    new_col = _require_new_column_name(df, f"{column}_was_missing")
    result = df.copy()
    flag = result[column].isna().astype(int)
    result[new_col] = flag
    return result, f"Added indicator column '{new_col}' ({int(flag.sum()):,} row(s) flagged)."


def clip_column_outliers(
    df: pd.DataFrame, column: str, method: str, low: float, high: float
) -> tuple[pd.DataFrame, str]:
    _require_column(df, column)
    if method not in CLIP_METHODS:
        raise ValueError(f"Unsupported clipping method: {method}")
    result = df.copy()
    numeric = pd.to_numeric(result[column], errors="coerce")
    if method == "Percentile":
        if not (0 <= low < high <= 100):
            raise ValueError("Percentile bounds must satisfy 0 <= low < high <= 100.")
        lower_bound, upper_bound = numeric.quantile(low / 100), numeric.quantile(high / 100)
    elif method == "Std devs from mean":
        mean, std = numeric.mean(), numeric.std()
        lower_bound, upper_bound = mean - low * std, mean + high * std
    else:  # Manual bounds
        if not low < high:
            raise ValueError("Lower bound must be less than upper bound.")
        lower_bound, upper_bound = low, high
    clipped = numeric.clip(lower=lower_bound, upper=upper_bound)
    changed = int(((numeric != clipped) & numeric.notna()).sum())
    result[column] = clipped
    return (
        result,
        f"Clipped '{column}' to [{lower_bound:.4g}, {upper_bound:.4g}] ({changed:,} value(s) changed).",
    )


def threshold_flag_column(
    df: pd.DataFrame, column: str, operator: str, threshold: float, new_column: str
) -> tuple[pd.DataFrame, str]:
    _require_column(df, column)
    if operator not in THRESHOLD_OPERATORS:
        raise ValueError(f"Unsupported comparison: {operator}")
    new_column = _require_new_column_name(df, new_column)
    result = df.copy()
    numeric = pd.to_numeric(result[column], errors="coerce")
    if operator == ">":
        mask = numeric > threshold
    elif operator == ">=":
        mask = numeric >= threshold
    elif operator == "<":
        mask = numeric < threshold
    elif operator == "<=":
        mask = numeric <= threshold
    elif operator == "==":
        mask = numeric == threshold
    else:  # "!="
        mask = numeric != threshold
    flag = mask.fillna(False).astype(int)
    result[new_column] = flag
    return (
        result,
        f"Created '{new_column}' = 1 where '{column}' {operator} {threshold:g} ({int(flag.sum()):,} row(s)).",
    )


def combine_columns(
    df: pd.DataFrame, left: str, operator: str, right: str, new_column: str
) -> tuple[pd.DataFrame, str]:
    _require_column(df, left)
    _require_column(df, right)
    if operator not in COMBINE_OPERATORS:
        raise ValueError(f"Unsupported operator: {operator}")
    new_column = _require_new_column_name(df, new_column)
    result = df.copy()
    left_series = pd.to_numeric(result[left], errors="coerce")
    right_series = pd.to_numeric(result[right], errors="coerce")
    with np.errstate(divide="ignore", invalid="ignore"):
        if operator == "+":
            combined = left_series + right_series
        elif operator == "-":
            combined = left_series - right_series
        elif operator == "*":
            combined = left_series * right_series
        else:  # "/"
            combined = left_series / right_series.replace(0, np.nan)
    result[new_column] = combined
    already_missing = int((left_series.isna() | right_series.isna()).sum())
    new_missing = int(combined.isna().sum()) - already_missing
    note = f" ({new_missing:,} divide-by-zero result(s) set to NaN)" if operator == "/" and new_missing > 0 else ""
    return result, f"Created '{new_column}' = '{left}' {operator} '{right}'{note}."


_TRANSFORM_DISPATCH: dict[str, Callable[..., tuple[pd.DataFrame, str]]] = {
    "rename_column": rename_column,
    "cast_column_dtype": cast_column_dtype,
    "drop_columns": drop_columns,
    "clean_text_column": clean_text_column,
    "extract_datetime_parts": extract_datetime_parts,
    "filter_rows": filter_rows,
    "dedupe_rows": dedupe_rows,
    "drop_missing_rows": drop_missing_rows,
    "fill_missing_column": fill_missing_column,
    "flag_missing_column": flag_missing_column,
    "clip_column_outliers": clip_column_outliers,
    "threshold_flag_column": threshold_flag_column,
    "combine_columns": combine_columns,
}


def apply_transform_step(df: pd.DataFrame, step: TransformStep) -> tuple[pd.DataFrame, str]:
    """Apply one pipeline step. `df` itself is never mutated -- each op
    function above works on its own copy."""
    handler = _TRANSFORM_DISPATCH.get(step.op)
    if handler is None:
        raise ValueError(f"Unknown data-preparation step: {step.op!r}")
    try:
        return handler(df, **step.params)
    except TypeError as exc:
        raise ValueError(f"Invalid parameters for step '{step.op}': {exc}") from exc


def apply_transform_pipeline(
    df: pd.DataFrame, steps: list[TransformStep]
) -> tuple[pd.DataFrame, list[str]]:
    """Fold apply_transform_step over `steps`, starting from a copy of `df`
    so the caller's frame is never mutated. On failure, raises a ValueError
    naming which step (1-based) and operation failed, and why -- so one bad
    step in a long pipeline is easy to find and fix."""
    current = df.copy()
    summaries: list[str] = []
    for index, step in enumerate(steps, start=1):
        label = TRANSFORM_OPS.get(step.op, TransformOpSpec(step.op, "")).label
        try:
            current, summary = apply_transform_step(current, step)
        except ValueError as exc:
            raise ValueError(f"Step {index} ({label}): {exc}") from exc
        summaries.append(summary)
    return current, summaries


def save_prepared_dataset(frame: pd.DataFrame, path: Path, cfg: RunConfig) -> None:
    """Preserve annotated spreadsheet options when editing its data rows."""
    save_dataset(frame, path)
    if cfg.input_mode != "Annotated spreadsheet":
        return
    from df_analyze.loading import load_spreadsheet

    metadata = load_spreadsheet(cfg.data_path, cfg.separator)[1]
    if not metadata.strip():
        return
    if path.suffix.lower() == ".csv":
        content = path.read_text(encoding="utf-8")
        path.write_text(metadata + "\n" + content, encoding="utf-8")
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path)
        sheet = workbook.active
        sheet.insert_rows(1)
        sheet.cell(1, 1, metadata)
        workbook.save(path)
        workbook.close()
    else:
        raise ValueError("Prepared annotated spreadsheets must remain CSV or XLSX.")


def prepare_external_test_files(cfg: RunConfig) -> list[str]:
    """Apply the same deterministic editing recipe and writer format to test files.

    validate_config disallows data-derived preparation statistics with external
    validation. Those belong in the analysis engine's fitted preprocessing.
    """
    prepared = []
    for i, source in enumerate(cfg.test_files):
        frame = load_full_dataset(Path(source), cfg.separator)
        if cfg.drop_duplicates:
            frame, _ = drop_duplicate_rows(frame)
        frame, _ = apply_transform_pipeline(frame, cfg.pipeline_steps)
        if frame.empty:
            raise ValueError(f"Preparation removed every row from test dataset {source}.")
        path = cfg.outdir / f"test_{i:02d}_prepared.parquet"
        save_dataset(frame, path)
        prepared.append(str(path))
    return prepared


def dataset_profile(df: pd.DataFrame) -> pd.DataFrame:
    """Summarize preview columns, including empty and nested JSON columns."""
    rows = []
    for index, column in enumerate(df.columns):
        values = df.iloc[:, index]
        missing = int(values.isna().sum())
        try:
            unique = int(values.nunique())
        except TypeError:
            unique = int(values.dropna().astype(str).nunique())
        rows.append(
            {
                "column": str(column),
                "dtype": str(values.dtype),
                "missing": missing,
                "missing_pct": 100.0 * missing / len(df) if len(df) else 0.0,
                "unique": unique,
            }
        )
    return pd.DataFrame(
        rows, columns=pd.Index(["column", "dtype", "missing", "missing_pct", "unique"])
    )


def list_runs(outdir: Path) -> list[Path]:
    """Find result directories beneath a workspace, or accept one directly."""
    candidates: list[tuple[float, Path]] = []
    for path in outdir.rglob("performance_long_table*.csv"):
        if not re.fullmatch(r"performance_long_table(?:_\d+)?\.csv", path.name):
            continue
        try:
            candidates.append((path.stat().st_mtime, path.parent))
        except OSError:
            continue  # A run may be removed while the explorer is refreshing.
    return [
        path for _, path in sorted(candidates, key=lambda item: item[0], reverse=True)
    ]


def find_latest_run(outdir: Path) -> Optional[Path]:
    runs = list_runs(outdir)
    return runs[0] if runs else None


def result_artifact(results_dir: Path, name: str) -> Path:
    """Find a standard artifact or the CLI's external-test numbered variant."""
    path = results_dir / name
    if path.exists():
        return path
    pattern = re.compile(re.escape(path.stem) + r"_\d+" + re.escape(path.suffix))
    candidates = sorted(p for p in results_dir.glob(f"{path.stem}_*{path.suffix}") if pattern.fullmatch(p.name))
    return candidates[0] if candidates else path


def run_artifact_root(results_dir: Path) -> Path:
    """Include preparation, tuning and analysis outputs for numbered runs too."""
    return results_dir.parent.parent if results_dir.parent.name == "results" else results_dir.parent


def report_artifact(results_dir: Path, target: str | None = None) -> Path:
    if target is None:
        return result_artifact(results_dir, "results_report.md")
    safe = re.sub(r"[^\w\.-]+", "_", target.strip()).strip("._") or "target"
    return result_artifact(results_dir, f"results_report_target_{safe}.md")


def load_results_table(results_dir: Path) -> pd.DataFrame:
    """Read performance_long_table.csv and add a display 'combo' column.

    NOTE: this file's column layout is inconsistent between df-analyze runs
    -- it sometimes has a stray leading unnamed index column and sometimes
    doesn't -- so we defensively drop anything named "Unnamed*" rather than
    assuming `index_col=0` (this broke the Streamlit app once already)."""
    long_table = pd.read_csv(result_artifact(results_dir, "performance_long_table.csv"))
    long_table = long_table.drop(
        columns=[c for c in long_table.columns if c.startswith("Unnamed")]
    )
    required = {"metric", "holdout", "model"}
    missing = required.difference(long_table.columns)
    if missing:
        raise ValueError(
            "Results table is missing columns: " + ", ".join(sorted(missing))
        )
    for column in ("selection", "embed_selector"):
        if column not in long_table:
            long_table[column] = "none"
        long_table[column] = long_table[column].fillna("none").astype(str)
    for column in ("holdout", "trainset", "5-fold"):
        if column not in long_table:
            long_table[column] = float("nan")
        long_table[column] = pd.to_numeric(long_table[column], errors="coerce")
    long_table["combo"] = (
        long_table["model"].fillna("unknown").astype(str)
        + " / "
        + long_table["selection"]
        + long_table["embed_selector"].apply(lambda s: f" ({s})" if s != "none" else "")
    )
    return long_table


def load_per_target_table(results_dir: Path) -> Optional[pd.DataFrame]:
    """Read performance_long_table_per_target.csv if this run used multiple
    targets (src/df_analyze/saving.py only writes this file in that case),
    building the same 'combo' display column as load_results_table but
    keeping the extra 'target' column. Returns None for a single-target run
    -- the file simply won't exist -- so callers can hide multi-target-only
    UI (like a target selector) rather than showing an error."""
    path = result_artifact(results_dir, "performance_long_table_per_target.csv")
    if not path.is_file():
        return None
    long_table = pd.read_csv(path)
    long_table = long_table.drop(
        columns=[c for c in long_table.columns if c.startswith("Unnamed")]
    )
    required = {"metric", "holdout", "model", "target"}
    missing = required.difference(long_table.columns)
    if missing:
        raise ValueError(
            "Per-target results table is missing columns: " + ", ".join(sorted(missing))
        )
    for column in ("selection", "embed_selector"):
        if column not in long_table:
            long_table[column] = "none"
        long_table[column] = long_table[column].fillna("none").astype(str)
    for column in ("holdout", "trainset", "5-fold"):
        if column not in long_table:
            long_table[column] = float("nan")
        long_table[column] = pd.to_numeric(long_table[column], errors="coerce")
    long_table["combo"] = (
        long_table["model"].fillna("unknown").astype(str)
        + " / "
        + long_table["selection"]
        + long_table["embed_selector"].apply(lambda s: f" ({s})" if s != "none" else "")
    )
    return long_table


def metric_options(long_table: pd.DataFrame) -> list[str]:
    return sorted(long_table["metric"].dropna().astype(str).unique().tolist())


def default_metric(metrics: list[str]) -> str:
    if not metrics:
        raise ValueError("No metrics available in this results table.")
    return "acc" if "acc" in metrics else metrics[0]


def ascending_for_metric(metric: str) -> bool:
    return metric in ERROR_METRICS


def filtered_sorted_view(long_table: pd.DataFrame, metric: str) -> pd.DataFrame:
    mask = (long_table["metric"] == metric) & long_table["holdout"].notna()
    view = long_table.loc[mask].copy()
    return view.sort_values("holdout", ascending=ascending_for_metric(metric))


def _validate_feature_count(raw: str, label: str, errors: list[str]) -> None:
    """Mirrors the CLI's own int_or_percent_parser: a non-negative number,
    where a value in [0, 1] means a fraction of all features and a value
    greater than 1 means an absolute feature count."""
    try:
        value = float(raw)
    except (TypeError, ValueError):
        errors.append(f"{label.capitalize()} must be a number (e.g. 0.5 or 20).")
        return
    if not math.isfinite(value) or value < 0:
        errors.append(f"{label.capitalize()} must be zero or greater.")


def validate_config(cfg: RunConfig, columns: Optional[list[str]] = None) -> list[str]:
    """Return actionable problems before starting an expensive analysis."""
    errors: list[str] = []
    try:
        validate_options(cfg.advanced_options)
    except ValueError as exc:
        errors.append(str(exc))
    if cfg.input_mode not in INPUT_MODES:
        errors.append("Choose a supported dataset input mode.")
    if cfg.input_mode == "Annotated spreadsheet" and cfg.data_path.suffix.lower() not in (".csv", ".xlsx"):
        errors.append("Annotated spreadsheets must be CSV or XLSX files.")
    if cfg.input_mode == "Training + test files" and not cfg.test_files:
        errors.append("Add at least one separate test dataset.")
    if cfg.input_mode == "Annotated spreadsheet" and cfg.test_files:
        errors.append("Separate test files cannot be combined with an annotated spreadsheet.")
    if cfg.tests_method not in ("list", "lodo"):
        errors.append("Choose list or leave-one-dataset-out evaluation.")
    if not isinstance(cfg.separator, str) or len(cfg.separator) != 1 or cfg.separator in ("\n", "\r"):
        errors.append("CSV delimiter must be one character (including tab), not a line break.")
    for path in cfg.test_files:
        if not Path(path).is_file() or Path(path).suffix.lower() not in SUPPORTED_SUFFIXES:
            errors.append(f"Choose an existing supported test dataset: {path}")
        if "," in path:
            errors.append("Test dataset paths cannot contain commas (CLI limitation).")
        if Path(path).resolve() == cfg.data_path.resolve():
            errors.append("A test dataset must be different from the training dataset.")
    if len(set(cfg.test_files)) != len(cfg.test_files):
        errors.append("Test datasets must not be repeated.")
    if cfg.test_files:
        for step in cfg.pipeline_steps:
            learns_fill = step.op == "fill_missing_column" and step.params.get("strategy") in ("Mean", "Median", "Mode")
            learns_clip = step.op == "clip_column_outliers" and step.params.get("method") != "Manual bounds"
            if learns_fill or learns_clip:
                errors.append("Separate test files require deterministic preparation steps. Use the analysis missing-value strategy for fitted imputation, or explicit constant fill/clipping bounds.")
                break
    if "--embed-select" in cfg.advanced_options and "embed" not in cfg.feat_select:
        errors.append("Enable embedded feature selection to use advanced embedded selectors.")
    if any(flag.startswith("--aer-") or flag == "--no-aer-smooth" for flag in cfg.advanced_options) and not cfg.adaptive_error:
        errors.append("Enable adaptive error scoring to use its advanced settings.")
    if any(flag.startswith("--aer-ens-") or flag == "--aer-ensemble-strategies" for flag in cfg.advanced_options) and not cfg.advanced_options.get("--aer-ensemble"):
        errors.append("Enable ensemble analysis to use ensemble settings.")
    if not cfg.data_path.is_file():
        errors.append("Choose an existing dataset file.")
    elif cfg.data_path.suffix.lower() not in SUPPORTED_SUFFIXES:
        errors.append("Choose a CSV, Parquet, XLSX or JSON dataset.")
    if not cfg.target.strip():
        errors.append("Choose a target column.")
    if cfg.mode not in ("classify", "regress"):
        errors.append("Task must be classify or regress.")
    choices = CLASSIFIERS if cfg.mode == "classify" else REGRESSORS
    if not cfg.models or set(cfg.models).difference(choices):
        errors.append("Select at least one supported model for this task.")
    metrics = CLS_METRICS if cfg.mode == "classify" else REG_METRICS
    if cfg.htune_metric not in metrics:
        errors.append("Choose a tuning metric that matches the task.")
    if not cfg.feat_select or set(cfg.feat_select).difference(FEAT_SELECT_METHODS):
        errors.append("Select a valid feature selection method.")
    if "none" in cfg.feat_select and len(cfg.feat_select) > 1:
        errors.append("Choose 'none' alone, or select feature selection methods.")
    if "embed" in cfg.feat_select and cfg.embed_model not in ("linear", "lgbm"):
        errors.append("Embedded selection requires a linear or LightGBM selector.")
    if "wrap" in cfg.feat_select:
        if cfg.wrapper_method not in ("step-up", "step-down"):
            errors.append("Choose step-up or step-down for wrapper selection.")
        if cfg.wrapper_model not in ("linear", "lgbm", "knn"):
            errors.append("Choose a model for wrapper selection.")
        _validate_feature_count(cfg.n_feat_wrapper, "wrapper feature count", errors)
        if not math.isfinite(cfg.redundant_threshold) or cfg.redundant_threshold < 0:
            errors.append("Redundancy threshold must be a non-negative number.")
        if not math.isfinite(cfg.redundant_corr_threshold) or cfg.redundant_corr_threshold < 0:
            errors.append("Redundancy correlation threshold must be a non-negative number.")
    if "filter" in cfg.feat_select:
        if cfg.filter_method not in FILTER_METHODS:
            errors.append("Choose a supported filter method.")
        if cfg.mode == "classify":
            if cfg.filter_assoc_cont_classify not in CONT_CLS_STATS:
                errors.append("Choose a supported continuous-feature association statistic.")
            if cfg.filter_assoc_cat_classify not in CAT_CLS_STATS:
                errors.append("Choose a supported categorical-feature association statistic.")
            if cfg.filter_pred_classify not in CLS_PRED_SCORES:
                errors.append("Choose a supported filter prediction score.")
        else:
            if cfg.filter_assoc_cont_regress not in CONT_REG_STATS:
                errors.append("Choose a supported continuous-feature association statistic.")
            if cfg.filter_assoc_cat_regress not in CAT_REG_STATS:
                errors.append("Choose a supported categorical-feature association statistic.")
            if cfg.filter_pred_regress not in REG_PRED_SCORES:
                errors.append("Choose a supported filter prediction score.")
        _validate_feature_count(cfg.n_feat_filter, "filter feature count", errors)
    if cfg.norm not in NORM_METHODS:
        errors.append("Choose a supported numeric scaling method.")
    if cfg.nan not in NAN_METHODS:
        errors.append("Choose a supported missing-value strategy.")
    if cfg.htune_trials < 1:
        errors.append("Tuning trials must be at least 1.")
    if not math.isfinite(cfg.test_val_size) or not 0 < cfg.test_val_size < 1:
        errors.append("Holdout fraction must be between 0 and 1 (exclusive).")
    if not 0 <= cfg.seed <= 2**32 - 1:
        errors.append("Seed must be between 0 and 4294967295.")
    if cfg.outdir.exists() and not cfg.outdir.is_dir():
        errors.append("Output location must be a directory.")
    roles = {
        "categorical": cfg.categoricals,
        "ordinal": cfg.ordinals,
        "excluded": cfg.drops,
    }
    used: dict[str, str] = {}
    for role, names in roles.items():
        for name in names:
            if name in used:
                errors.append(f"Column '{name}' has conflicting or repeated roles.")
            used[name] = role
            if "," in name:
                errors.append(
                    f"Rename '{name}': commas are not supported in column roles."
                )
    if cfg.target in used or cfg.target == cfg.grouper:
        errors.append(
            "The target cannot also be a feature role, excluded column, or group."
        )
    if cfg.grouper and cfg.grouper in used:
        errors.append("The grouping column must have its own role.")
    if columns is not None:
        referenced = [cfg.target, *used, *([cfg.grouper] if cfg.grouper else [])]
        missing = sorted(set(referenced).difference(columns))
        if missing:
            errors.append("Columns not found in this dataset: " + ", ".join(missing))
        if len(columns) != len(set(columns)):
            errors.append("Dataset column names must be unique.")
        if not set(columns).difference([cfg.target, *cfg.drops, cfg.grouper]):
            errors.append("Keep at least one predictor column.")
    if cfg.adaptive_error and cfg.mode != "classify":
        errors.append("Adaptive error analysis is available for classification.")
    if cfg.extra_targets:
        if set(cfg.extra_targets).intersection([*used, cfg.grouper]):
            errors.append("Additional targets cannot also be feature roles, excluded columns, or groups.")
        if columns is not None and not set(columns).difference([cfg.target, *cfg.extra_targets, *cfg.drops, cfg.grouper]):
            errors.append("Keep at least one predictor column outside the targets.")
        if cfg.target in cfg.extra_targets:
            errors.append("The primary target cannot also be listed as an extra target.")
        if len(cfg.extra_targets) != len(set(cfg.extra_targets)):
            errors.append("Extra targets must not repeat the same column.")
        if columns is not None:
            missing_targets = sorted(set(cfg.extra_targets).difference(columns))
            if missing_targets:
                errors.append("Extra target column(s) not found: " + ", ".join(missing_targets))
        if cfg.mt_agg_strategy not in MT_AGG_STRATEGIES:
            errors.append("Choose a supported multi-target aggregation strategy.")
        if cfg.mt_top_k.strip():
            try:
                top_k = int(cfg.mt_top_k)
            except ValueError:
                errors.append("Multi-target top-K must be a whole number.")
            else:
                if top_k < 1:
                    errors.append("Multi-target top-K must be at least 1.")
    for step in cfg.pipeline_steps:
        if step.op not in TRANSFORM_OPS:
            errors.append(f"Unknown data-preparation step: {step.op!r}")
            break
    try:
        parse_extra_args(cfg.extra_args)
    except ValueError as exc:
        errors.append(f"Extra arguments contain invalid quoting: {exc}")
    return list(dict.fromkeys(errors))


def apply_preset(cfg: RunConfig, name: str) -> RunConfig:
    """Return an independent configuration; preserve data and column roles."""
    if name not in PRESET_NAMES:
        raise ValueError(f"Unknown preset: {name}")
    result = config_from_json(config_to_json(cfg))
    linear = "lr" if cfg.mode == "classify" else "elastic"
    models, selection, trials, no_preds = {
        "Quick baseline": (["dummy", linear], ["none"], 10, True),
        "Balanced": (["dummy", linear, "lgbm"], ["filter", "embed"], 50, False),
        "Thorough": (
            ["dummy", linear, "lgbm", "rf", "catboost"],
            ["filter", "embed"],
            150,
            False,
        ),
    }[name]
    return replace(
        result,
        models=models,
        feat_select=selection,
        htune_trials=trials,
        no_preds=no_preds,
        embed_model="linear",
        norm="robust",
        nan="median",
        htune_metric="acc" if cfg.mode == "classify" else "mae",
    )


def config_to_json(cfg: RunConfig) -> str:
    data = asdict(cfg)
    data["data_path"], data["outdir"] = str(cfg.data_path), str(cfg.outdir)
    return json.dumps({"version": 1, "config": data}, indent=2, ensure_ascii=False)


def _steps_from_json(raw: object) -> list[TransformStep]:
    """Reconstruct and validate a list of TransformStep from parsed JSON,
    rejecting anything that isn't the plain {"op": ..., "params": {...}}
    shape with JSON-safe param values -- the same "no code, no callables"
    stance the rest of this file's JSON loading already takes."""
    if not isinstance(raw, list):
        raise ValueError("pipeline_steps must be a list.")
    steps: list[TransformStep] = []
    for item in raw:
        if not isinstance(item, dict) or "op" not in item or set(item) - {"op", "params"}:
            raise ValueError(
                "Each pipeline step must be an object with only 'op' and optional 'params'."
            )
        op = item["op"]
        if not isinstance(op, str) or op not in TRANSFORM_OPS:
            raise ValueError(f"Unknown data-preparation step: {op!r}")
        params = item.get("params", {})
        if not isinstance(params, dict):
            raise ValueError(f"Parameters for step '{op}' must be an object.")
        for key, value in params.items():
            if not isinstance(key, str):
                raise ValueError(f"Parameter names for step '{op}' must be text.")
            if isinstance(value, list):
                if not all(isinstance(v, str) for v in value):
                    raise ValueError(
                        f"List parameter '{key}' for step '{op}' must contain only text."
                    )
            elif value is not None and not isinstance(value, (str, int, float, bool)):
                raise ValueError(f"Parameter '{key}' for step '{op}' has an unsupported value type.")
        steps.append(TransformStep(op=op, params=dict(params)))
    return steps


def config_from_json(text: str) -> RunConfig:
    """Load plain JSON only; reject malformed settings without executing code."""
    try:
        document = json.loads(text)
        if not isinstance(document, dict) or document.get("version", 1) != 1:
            raise ValueError("Unsupported configuration format or version.")
        data = document.get("config", document)
        if not isinstance(data, dict):
            raise ValueError("Configuration must be a JSON object.")
        data = dict(data)
        valid_keys = {item.name for item in fields(RunConfig)}
        if set(data).difference(valid_keys):
            raise ValueError("Configuration contains unrecognized settings.")
        for key in ("data_path", "outdir"):
            if key in data:
                if not isinstance(data[key], str) or not data[key].strip():
                    raise ValueError(f"{key} must be a nonempty path string.")
                data[key] = Path(data[key]).expanduser()
        list_keys = ("models", "feat_select", "categoricals", "ordinals", "drops", "extra_targets", "test_files")
        for key in list_keys:
            if key in data and (
                not isinstance(data[key], list)
                or not all(isinstance(item, str) for item in data[key])
            ):
                raise ValueError(f"{key} must be a list of column names or choices.")
        for key in ("htune_trials", "seed"):
            if key in data and type(data[key]) is not int:
                raise ValueError(f"{key} must be an integer.")
        bool_keys = ("no_preds", "adaptive_error", "drop_duplicates", "redundant_wrapper_selection")
        for key in bool_keys:
            if key in data and type(data[key]) is not bool:
                raise ValueError(f"{key} must be true or false.")
        float_keys = ("test_val_size", "redundant_threshold", "redundant_corr_threshold")
        for key in float_keys:
            if key in data and (
                type(data[key]) not in (float, int) or not math.isfinite(data[key])
            ):
                raise ValueError(f"{key} must be a finite number.")
        if "pipeline_steps" in data:
            data["pipeline_steps"] = _steps_from_json(data["pipeline_steps"])
        if "advanced_options" in data:
            data["advanced_options"] = validate_options(data["advanced_options"])
        for key in valid_keys.difference(
            [
                *list_keys,
                *bool_keys,
                *float_keys,
                "data_path",
                "outdir",
                "htune_trials",
                "seed",
                "pipeline_steps",
                "advanced_options",
            ]
        ):
            if key in data and not (key == "grouper" and data[key] is None):
                if not isinstance(data[key], str):
                    raise ValueError(f"{key} must be text.")
        return RunConfig(**data)
    except (TypeError, KeyError, json.JSONDecodeError) as exc:
        raise ValueError(f"Invalid experiment configuration: {exc}") from exc
